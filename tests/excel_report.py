import sys; from pathlib import Path; sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
from pathlib import Path
from openpyxl import load_workbook
from analytics.visualization.excel.workbook import build, OUTPUT

build()
assert Path(OUTPUT).exists()
assert Path(OUTPUT).stat().st_size > 0
wb = load_workbook(OUTPUT, read_only=True)
required = {"Executive Summary","Certified Data","Market Analysis","Competitive Intelligence","Customer Analytics","Trends","Forecast","Opportunity & Risk","Relationships","Composition","Evidence","Raw Data","Visualization Decisions","README"}
assert required.issubset(set(wb.sheetnames))
wb.close()
print("EXCEL REPORT V2 TEST : PASS")
